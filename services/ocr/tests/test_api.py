"""Integração com PostgreSQL/PostGIS: API, persistência, lotes e exportação."""
import io
import zipfile

import openpyxl
import pytest

import fabrica as F

pytestmark = pytest.mark.db


def _enviar(cliente, conteudo, nome="req.pdf", ct="application/pdf"):
    return cliente.post("/documentos", files={"arquivo": (nome, conteudo, ct)})


def _detalhe(cliente, uid):
    r = cliente.get(f"/documentos/{uid}")
    assert r.status_code == 200, r.text
    return r.json()


def _campo(det, nome, indice=None):
    return next(c for c in det["requerimento"]["campos"] if c["campo"] == nome and c["indice"] == indice)


def test_ct40_upload_pdf_real(cliente, pdf_exemplo):
    r = _enviar(cliente, pdf_exemplo, "ramon-teste.pdf")
    assert r.status_code == 201, r.text
    d = _detalhe(cliente, r.json()["uuid"])
    assert d["status"] == "revisao"
    req = d["requerimento"]
    assert req["cpf_cnpj"] == "059.749.653-61" and req["pessoa_id"] == 1
    assert req["municipio"] == "Altos" and req["municipio_id"] == 1
    assert req["procedimento_id"] == 1 and req["tipo_intervencao_id"] == 1
    assert [i["cultura"] for i in req["irrigacao"]] == ["Alface", "Acerola", "Abacaxi (solo limpo)"]
    assert all(i["sistema_id"] for i in req["irrigacao"])
    assert req["finalidades"][0]["tipo_finalidade_id"] == 1
    assert req["pendencias"] == ["Latitude: segundos (91.06) fora do intervalo 0–59,99"]
    assert _campo(d, "latitude")["bbox"] and _campo(d, "latitude")["pagina"] == 1
    assert d["paginas_lidas"][0]["metodo"] == "texto_nativo"


def test_ct41_duplicado(cliente, pdf_exemplo):
    assert _enviar(cliente, pdf_exemplo).status_code == 201
    r = _enviar(cliente, pdf_exemplo, "copia.pdf")
    assert r.status_code == 200 and r.json()["duplicado"] is True


def test_ct42_corrigir_latitude(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"latitude": "-5°02'41.06\""}},
                      headers={"X-Usuario-Id": "42"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extraido", d["requerimento"]["pendencias"]
    lat = _campo(d, "latitude")
    assert lat["corrigido_manual"] and lat["valido"] and lat["confianca"] == 100
    assert lat["bbox"]  # mantém a posição original para a tela de revisão
    assert d["requerimento"]["latitude"] == pytest.approx(-5.0447389, abs=1e-6)
    hist = cliente.get(f"/documentos/{uid}/revisoes").json()
    assert hist[0]["campo"] == "latitude" and hist[0]["valor_anterior"] == "-5°02'91.06\"" and hist[0]["usuario_id"] == 42
    # validar agora funciona
    assert cliente.post(f"/documentos/{uid}/validar").json()["status"] == "validado"


def test_ct43_validar_com_pendencias(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    r = cliente.post(f"/documentos/{uid}/validar")
    assert r.status_code == 409


def test_ct44_validar_forcado_com_justificativa(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    assert cliente.post(f"/documentos/{uid}/validar", json={"forcar": True}).status_code == 409
    r = cliente.post(f"/documentos/{uid}/validar", json={"forcar": True, "justificativa": "conferido em campo"})
    assert r.status_code == 200 and r.json()["status"] == "validado"
    hist = cliente.get(f"/documentos/{uid}/revisoes").json()
    assert hist[-1]["campo"] == "__validacao_com_pendencias__" and hist[-1]["valor_novo"] == "conferido em campo"


def test_ct45_ponto_fora_do_municipio(cliente):
    # polígono fictício de Altos no stub vai de -42,75 a -42,25 de longitude
    pdf = F.pdf_requerimento(F.DadosReq(longitude="-42°10'00\""))
    d = _detalhe(cliente, _enviar(cliente, pdf).json()["uuid"])
    assert any("fora do município declarado" in p for p in d["requerimento"]["pendencias"])
    pdf_ok = F.pdf_requerimento(F.DadosReq(longitude="-42°30'00\""))
    d2 = _detalhe(cliente, _enviar(cliente, pdf_ok, "ok.pdf").json()["uuid"])
    assert d2["status"] == "extraido", d2["requerimento"]["pendencias"]


def test_ct46_nome_divergente_do_cadastro(cliente):
    pdf = F.pdf_requerimento(F.DadosReq(nome="José Pereira Lima", documento="05974965361"))
    d = _detalhe(cliente, _enviar(cliente, pdf).json()["uuid"])
    nome = _campo(d, "requerente_nome")
    assert "Ramon Teste" in nome["mensagem"] and d["status"] == "revisao"


def test_ct47_pdf_com_tres_requerimentos(cliente):
    pdf = F.pdf_requerimento(F.DadosReq(), F.DadosReq(localidade="Sítio Novo"), F.DadosReq(municipio="Teresina",
                             longitude="-42°50'00\""))
    r = _enviar(cliente, pdf, "malote.pdf")
    pai = _detalhe(cliente, r.json()["uuid"])
    assert pai["status"] == "dividido" and len(pai["filhos"]) == 3
    filhos = [_detalhe(cliente, f["uuid"]) for f in pai["filhos"]]
    assert [f["requerimento"]["localidade"] for f in filhos] == ["Povoado Boa Esperança", "Sítio Novo",
                                                                  "Povoado Boa Esperança"]
    assert filhos[2]["requerimento"]["municipio"] == "Teresina"
    assert all(f["documento_pai_uuid"] == pai["uuid"] for f in filhos)
    assert cliente.get(f"/documentos/{filhos[1]['uuid']}/arquivo").content.startswith(b"%PDF")


def _zip(arquivos):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for nome, dados in arquivos:
            z.writestr(nome, dados)
    return buf.getvalue()


def test_ct48_lote_zip_misto(cliente, pdf_exemplo):
    img = F.png(F.renderizar(F.pdf_requerimento(F.DadosReq(localidade="Lagoa Seca"))))
    z = _zip([
        ("a/requerimento1.pdf", pdf_exemplo),
        ("requerimento2.pdf", F.pdf_requerimento(F.DadosReq(localidade="Baixão"))),
        ("foto.png", img),
        ("leia-me.txt", b"nada aqui"),
        ("copia.pdf", pdf_exemplo),
        ("__MACOSX/._lixo", b"x"),
    ])
    r = cliente.post("/lotes", files=[("arquivos", ("malote.zip", z, "application/zip"))], data={"nome": "Malote 1"})
    assert r.status_code == 201, r.text
    corpo = r.json()
    assert corpo["recebidos"] == 5 and corpo["novos"] == 4
    assert [d["arquivo"] for d in corpo["duplicados"]] == ["copia.pdf"]
    assert corpo["formato_nao_suportado"] == ["leia-me.txt"]
    lote = cliente.get(f"/lotes/{corpo['uuid']}").json()
    assert lote["total"] == 4 and lote["processados"] == 4 and lote["com_erro"] == 1
    assert lote["status"] == "concluido_com_erros"
    erro = next(d for d in lote["documentos"] if d["status"] == "erro")
    assert "formato não suportado" in erro["erro"]


def test_ct49_reprocessar_erros(cliente):
    z = _zip([("ok.pdf", F.pdf_requerimento(F.DadosReq())), ("ruim.txt", b"x")])
    lote = cliente.post("/lotes", files=[("arquivos", ("l.zip", z, "application/zip"))]).json()
    r = cliente.post(f"/lotes/{lote['uuid']}/reprocessar-erros")
    assert r.json()["reenfileirados"] == 1
    docs = cliente.get(f"/lotes/{lote['uuid']}").json()["documentos"]
    ok = next(d for d in docs if d["nome_original"] == "ok.pdf")
    assert ok["status"] == "extraido"


def test_ct50_exportacao(cliente, pdf_exemplo):
    lote = cliente.post("/lotes", files=[
        ("arquivos", ("a.pdf", pdf_exemplo, "application/pdf")),
        ("arquivos", ("b.pdf", F.pdf_requerimento(F.DadosReq()), "application/pdf")),
    ]).json()
    csv = cliente.get(f"/lotes/{lote['uuid']}/exportar?formato=csv")
    assert csv.status_code == 200 and csv.content.startswith("\ufeff".encode())
    texto = csv.content.decode("utf-8-sig")
    linhas = texto.strip().splitlines()
    assert len(linhas) == 3 and linhas[0].startswith("Documento (uuid);Arquivo;Status")
    assert "Regularização de Poço Tubular" in texto and "-42,7089944" in texto
    x = cliente.get(f"/lotes/{lote['uuid']}/exportar?formato=xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Requerimentos", "Irrigação", "Detalhes das finalidades"]
    assert wb["Irrigação"].max_row == 7  # cabeçalho + 3 + 3


def test_ct51_api_key(cliente, monkeypatch):
    from ocr_requerimentos.config import config

    monkeypatch.setattr(config(), "api_key", "segredo")
    assert cliente.get("/lotes").status_code == 401
    assert cliente.get("/lotes", headers={"X-API-Key": "segredo"}).status_code == 200
    assert cliente.get("/saude").status_code == 200  # saúde fica aberta para o healthcheck


def test_ct52_corrigir_irrigacao(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    novo = [
        {"sistema": "Aspersão por canhão hidráulico", "area": "10", "cultura": "Alface"},
        {"sistema": "Aspersão por sistema autopropelido", "area": "3,5", "cultura": "Goiaba"},
    ]
    d = cliente.patch(f"/documentos/{uid}/correcoes", json={"irrigacao": novo}).json()
    itens = d["requerimento"]["irrigacao"]
    assert [(i["cultura"], i["area_ha"]) for i in itens] == [("Alface", 10.0), ("Goiaba", 3.5)]
    assert not _campo(d, "irrigacao_cultura", 1)["corrigido_manual"]
    assert _campo(d, "irrigacao_cultura", 2)["corrigido_manual"]
    assert _campo(d, "irrigacao_area", 2)["corrigido_manual"]
    assert _campo(d, "irrigacao_sistema", 1)["bbox"]  # linha mantida conserva a caixa


def test_ct53_imagem_da_pagina(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    r = cliente.get(f"/documentos/{uid}/paginas/1/imagem")
    assert r.status_code == 200 and r.headers["content-type"] == "image/jpeg" and r.content[:2] == b"\xff\xd8"
    assert cliente.get(f"/documentos/{uid}/paginas/9/imagem").status_code == 404


def test_ct54_reprocessar_validado(cliente):
    uid = _enviar(cliente, F.pdf_requerimento(F.DadosReq())).json()["uuid"]
    assert cliente.post(f"/documentos/{uid}/validar").status_code == 200
    assert cliente.post(f"/documentos/{uid}/reprocessar").status_code == 409


def test_fila_de_revisao_e_busca(cliente, pdf_exemplo):
    _enviar(cliente, pdf_exemplo)
    _enviar(cliente, F.pdf_requerimento(F.DadosReq()), "ok.pdf")
    fila = cliente.get("/revisao").json()
    assert fila["total"] == 1 and fila["itens"][0]["requerente_nome"] == "ramon teste"
    busca = cliente.get("/documentos", params={"busca": "Maria"}).json()
    assert busca["total"] == 1
    # CPF 059.749.653-61 com ou sem máscara, inteiro ou em parte
    for termo in ("05974965361", "059.749.653-61", "059749653-61", "059.749653.61", "749653", " 653-61 "):
        r = cliente.get("/documentos", params={"busca": termo}).json()
        assert [d["requerente_nome"] for d in r["itens"]] == ["ramon teste"], termo
    assert cliente.get("/documentos", params={"busca": "99999999999"}).json()["total"] == 0
    assert cliente.get("/estatisticas").json()["por_status"] == {"revisao": 1, "extraido": 1}


def test_referencias(cliente):
    municipios = cliente.get("/referencias/municipios").json()
    assert {"id": 1, "nome": "Altos"} in municipios and all(m["nome"] != "Timon" for m in municipios)
    assert len(cliente.get("/referencias/culturas").json()) == 206  # lista do formulário
    assert cliente.get("/referencias/xyz").status_code == 404


def test_upload_zip_avulso_rejeitado(cliente):
    assert _enviar(cliente, b"PK\x03\x04", "a.zip", "application/zip").status_code == 400


def test_uuid_invalido(cliente):
    assert cliente.get("/documentos/nao-e-uuid").status_code == 404


# ------------------------------------------------------------------ CT-86..89: detalhes das finalidades no banco
FORM = F.__file__.rsplit("/", 1)[0] + "/fixtures/formulario/"


def _pdf_form(nome):
    return open(FORM + f"{nome}.pdf", "rb").read()


def _fin(d, chave):
    return next(f for f in d["requerimento"]["finalidades"] if f["chave"] == chave)


@pytest.mark.parametrize("nome", ["todas_1", "todas_2"])
def test_ct86_detalhes_gravados(cliente, banco, nome):
    from sqlalchemy import text

    uid = _enviar(cliente, _pdf_form(nome), f"{nome}.pdf").json()["uuid"]
    d = _detalhe(cliente, uid)
    assert d["status"] == "extraido", d["requerimento"]["pendencias"]
    fins = d["requerimento"]["finalidades"]
    assert len(fins) == 10 and all(f["tipo_finalidade_id"] for f in fins)  # todas ligadas ao SIGA
    assert all(f["estrutura"]["nome"] == f["texto"] for f in fins)
    # cada detalhe também vira campo_extraido, com posição na página para a revisão
    cab = [c for c in d["requerimento"]["campos"] if c["campo"] == "criacao_animal.cabecas"]
    assert [c["indice"] for c in cab] == [1, 2, 3][: len(cab)] and all(c["bbox"] for c in cab)
    assert _campo(d, "termoeletrica.potencia")["rotulo"] == "Termoelétrica — Potência Instalada (MW)"
    with banco.connect() as c:
        chaves = c.execute(text(
            "SELECT f.chave, f.detalhes FROM ocr.requerimento_finalidade f "
            "JOIN ocr.requerimento r ON r.id = f.requerimento_id JOIN ocr.documento d ON d.id = r.documento_id "
            "WHERE d.uuid = CAST(:u AS uuid) ORDER BY f.id"), {"u": uid}).all()
    assert [k for k, _ in chaves][:2] == ["abastecimento_publico", "aquicultura_tanque"]
    assert dict(chaves)["termoeletrica"]["combustivel"]


@pytest.mark.parametrize("nome", ["fin04_2", "fin03_2"])
def test_ct87_corrigir_item_de_tabela(cliente, nome):
    chave, coluna = {"fin04_2": ("criacao_animal", "cabecas"), "fin03_2": ("consumo_humano", "pessoas")}[nome]
    uid = _enviar(cliente, _pdf_form(nome), f"{nome}.pdf").json()["uuid"]
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"detalhes": {f"{chave}.{coluna}#2": "150"}},
                      headers={"X-Usuario-Id": "7"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert _fin(d, chave)["detalhes"]["itens"][1][coluna] == 150
    assert _campo(d, f"{chave}.{coluna}", 2)["corrigido_manual"]
    assert d["status"] == "extraido"
    hist = cliente.get(f"/documentos/{uid}/revisoes").json()
    assert hist[-1]["campo"] == f"{chave}.{coluna}#2" and hist[-1]["valor_novo"] == "150"
    # remover o item 1 (todas as colunas vazias): os seguintes sobem e a correção acompanha
    colunas = [c["chave"] for c in _fin(d, chave)["estrutura"]["colunas"]]
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"detalhes": {f"{chave}.{c}#1": "" for c in colunas}})
    d = r.json()
    itens = _fin(d, chave)["detalhes"]["itens"]
    assert len(itens) == 2 and itens[0][coluna] == 150
    assert _campo(d, f"{chave}.{coluna}", 1)["corrigido_manual"]


@pytest.mark.parametrize("nome", ["incompleto_1", "incompleto_2"])
def test_ct88_completar_finalidade_vazia(cliente, nome):
    uid = _enviar(cliente, _pdf_form(nome), f"{nome}.pdf").json()["uuid"]
    d = _detalhe(cliente, uid)
    assert d["status"] == "revisao"
    preencher = {
        "incompleto_1": {"criacao_animal.especie#1": "Bovino de Corte", "criacao_animal.sistema#1": "Extensivo",
                         "criacao_animal.cabecas#1": "40"},
        "incompleto_2": {"abastecimento_publico.prestador": "Administração direta (prefeitura)",
                         "abastecimento_publico.municipio#1": "URUÇUÍ", "abastecimento_publico.localidade#1": "Sede",
                         "abastecimento_publico.habitantes#1": "2000",
                         "aquicultura_tanque.profundidade": "2", "aquicultura_tanque.area_tanques": "300",
                         "aquicultura_tanque.especie#1": "Piscicultura Tambaqui",
                         "termoeletrica.combustivel": "Biomassa", "termoeletrica.potencia": "5",
                         "termoeletrica.producao_energia": "900"},
    }[nome]
    corpo = {"detalhes": preencher}
    if nome == "incompleto_1":  # irrigação continua na tabela própria
        corpo["irrigacao"] = [{"sistema": "Gotejamento", "area": "3", "cultura": "Tomate"}]
    r = cliente.patch(f"/documentos/{uid}/correcoes", json=corpo)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extraido", d["requerimento"]["pendencias"]


def test_ct89_detalhe_invalido(cliente):
    uid = _enviar(cliente, _pdf_form("fin10_1"), "fin10_1.pdf").json()["uuid"]
    for chave, msg in (("termoeletrica.inexistente", "desconhecido"), ("criacao_animal.cabecas", "item"),
                       ("termoeletrica.potencia#1", "campo sem item"), ("irrigacao.area#1", "desconhecido")):
        r = cliente.patch(f"/documentos/{uid}/correcoes", json={"detalhes": {chave: "1"}})
        assert r.status_code == 422 and msg in r.json()["detail"], (chave, r.text)
    # valor fora da lista do formulário vira pendência, não erro
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"detalhes": {"termoeletrica.combustivel": "Lenha molhada"}})
    d = r.json()
    assert d["status"] == "revisao"
    assert any(p.startswith("Termoelétrica — Combustível Principal") for p in d["requerimento"]["pendencias"])


# ------------------------------------------------------------------ CT-90..97: integridade e rastreabilidade
def test_ct90_reprocessar_mantem_correcoes(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    novo = [{"sistema": "Gotejamento", "area": "4", "cultura": "Tomate"}]
    d = cliente.patch(f"/documentos/{uid}/correcoes",
                      json={"campos": {"latitude": "-5°02'41.06\""}, "irrigacao": novo}).json()
    assert d["status"] == "extraido"
    assert cliente.post(f"/documentos/{uid}/reprocessar").status_code == 200
    d = _detalhe(cliente, uid)
    assert d["status"] == "extraido", d["requerimento"]["pendencias"]
    lat = _campo(d, "latitude")
    assert lat["corrigido_manual"] and lat["valor_bruto"] == "-5°02'41.06\""
    assert [(i["cultura"], i["area_ha"]) for i in d["requerimento"]["irrigacao"]] == [("Tomate", 4.0)]
    hist = cliente.get(f"/documentos/{uid}/revisoes").json()
    assert hist[-1]["campo"] == "__reprocessamento__" and hist[-1]["valor_novo"].endswith("após nova leitura")
    assert sum(h["campo"] == "latitude" for h in hist) == 1  # reaplicar não duplica a auditoria


def test_ct91_reprocessar_sem_correcoes_nao_audita(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    cliente.post(f"/documentos/{uid}/reprocessar")
    assert cliente.get(f"/documentos/{uid}/revisoes").json() == []
    assert _detalhe(cliente, uid)["status"] == "revisao"


def test_ct92_mesmo_arquivo_em_dois_lotes(cliente, pdf_exemplo):
    a = cliente.post("/lotes", files=[("arquivos", ("a.pdf", pdf_exemplo, "application/pdf"))]).json()
    b = cliente.post("/lotes", files=[("arquivos", ("reenvio.pdf", pdf_exemplo, "application/pdf"))]).json()
    assert b["novos"] == 0 and [d["arquivo"] for d in b["duplicados"]] == ["reenvio.pdf"]
    la, lb = cliente.get(f"/lotes/{a['uuid']}").json(), cliente.get(f"/lotes/{b['uuid']}").json()
    assert la["total"] == lb["total"] == 1 and lb["em_revisao"] == 1 and lb["status"] == "concluido"
    assert la["documentos"][0]["uuid"] == lb["documentos"][0]["uuid"]
    assert lb["documentos"][0]["lote_uuid"] == a["uuid"]  # lote de origem
    linhas = cliente.get(f"/lotes/{b['uuid']}/exportar").content.decode("utf-8-sig").strip().splitlines()
    assert len(linhas) == 2
    # validar o documento atualiza os dois lotes
    uid = la["documentos"][0]["uuid"]
    cliente.post(f"/documentos/{uid}/validar", json={"forcar": True, "justificativa": "ok"})
    assert cliente.get(f"/lotes/{a['uuid']}").json()["validados"] == 1
    assert cliente.get(f"/lotes/{b['uuid']}").json()["validados"] == 1


def test_ct93_partes_de_pdf_dividido_nos_lotes(cliente):
    pdf = F.pdf_requerimento(F.DadosReq(), F.DadosReq(localidade="Sítio Novo"))
    a = cliente.post("/lotes", files=[("arquivos", ("malote.pdf", pdf, "application/pdf"))]).json()
    b = cliente.post("/lotes", files=[("arquivos", ("malote2.pdf", pdf, "application/pdf"))]).json()
    for lote in (a, b):
        corpo = cliente.get(f"/lotes/{lote['uuid']}").json()
        assert corpo["total"] == 2 and len(corpo["documentos"]) == 3  # original dividido + 2 partes


def test_ct94_versao_concorrente(cliente, pdf_exemplo):
    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    v = _detalhe(cliente, uid)["requerimento"]["versao"]
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"localidade": "Sede"}, "versao": v})
    assert r.status_code == 200 and r.json()["requerimento"]["versao"] > v
    # segunda pessoa ainda com a versão antiga
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"localidade": "Outra"}, "versao": v})
    assert r.status_code == 409 and "recarregue" in r.json()["detail"]
    r = cliente.post(f"/documentos/{uid}/validar", json={"forcar": True, "justificativa": "x", "versao": v})
    assert r.status_code == 409
    assert _detalhe(cliente, uid)["requerimento"]["localidade"] == "Sede"
    # sem versão (clientes antigos) continua aceito
    assert cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"localidade": "Outra"}}).status_code == 200


def test_ct95_sem_edicao_durante_leitura(cliente, banco, pdf_exemplo):
    from sqlalchemy import text

    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    with banco.begin() as c:
        c.execute(text("UPDATE ocr.documento SET status = 'processando' WHERE uuid = CAST(:u AS uuid)"), {"u": uid})
    assert cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"localidade": "x"}}).status_code == 409
    assert cliente.post(f"/documentos/{uid}/validar").status_code == 409


def test_ct96_restricoes_do_banco(cliente, banco, pdf_exemplo):
    import sqlalchemy as sa

    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    invalidos = [
        "UPDATE ocr.documento SET status = 'qualquer' WHERE uuid = CAST(:u AS uuid)",
        "UPDATE ocr.requerimento SET confianca_geral = 101 WHERE documento_id = "
        "(SELECT id FROM ocr.documento WHERE uuid = CAST(:u AS uuid))",
        "UPDATE ocr.requerimento SET cpf_cnpj = '059.749.653-6l' WHERE documento_id = "
        "(SELECT id FROM ocr.documento WHERE uuid = CAST(:u AS uuid))",
        "UPDATE ocr.requerimento SET status_validacao = 'validado', validado_em = NULL WHERE documento_id = "
        "(SELECT id FROM ocr.documento WHERE uuid = CAST(:u AS uuid))",
        "INSERT INTO ocr.campo_extraido (requerimento_id, campo, obrigatorio, corrigido_manual) "
        "SELECT r.id, 'latitude', false, false FROM ocr.requerimento r "
        "JOIN ocr.documento d ON d.id = r.documento_id WHERE d.uuid = CAST(:u AS uuid)",
    ]
    for sql in invalidos:
        with pytest.raises(sa.exc.IntegrityError), banco.begin() as c:
            c.execute(sa.text(sql), {"u": uid})
    # o mesmo CPF pode ter vários requerimentos
    _enviar(cliente, F.pdf_requerimento(F.DadosReq(documento="05974965361", localidade="Outro")), "outro.pdf")
    with banco.connect() as c:
        assert c.scalar(sa.text("SELECT count(*) FROM ocr.requerimento WHERE cpf_cnpj = '059.749.653-61'")) == 2


def test_ct97_finalidade_repetida_na_correcao(cliente):
    uid = _enviar(cliente, _pdf_form("fin04_2"), "fin04_2.pdf").json()["uuid"]
    textos = [f["texto"] for f in _detalhe(cliente, uid)["requerimento"]["finalidades"]]
    r = cliente.patch(f"/documentos/{uid}/correcoes", json={"finalidades": textos + textos[:1]})
    assert r.status_code == 200, r.text
    assert len(r.json()["requerimento"]["finalidades"]) == len(textos) + 1


def test_ct98_acuracia_e_versao_do_pipeline(cliente, banco, pdf_exemplo):
    from sqlalchemy import text

    from ocr_requerimentos.pipeline import VERSAO_PIPELINE

    uid = _enviar(cliente, pdf_exemplo).json()["uuid"]
    antes = _campo(_detalhe(cliente, uid), "latitude")
    cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"latitude": "-5°02'41.06\""}})
    cliente.patch(f"/documentos/{uid}/correcoes", json={"campos": {"localidade": "Sede"}})  # nova gravação
    cliente.post(f"/documentos/{uid}/validar")
    with banco.connect() as c:
        versao, cfg = c.execute(text("SELECT versao_pipeline, config_extracao FROM ocr.documento "
                                     "WHERE uuid = CAST(:u AS uuid)"), {"u": uid}).one()
        assert versao == VERSAO_PIPELINE and cfg["motor_ocr"] and cfg["limiar_confianca"] == 80.0
        ocr, conf = c.execute(text(
            "SELECT valor_ocr, confianca_ocr FROM ocr.campo_extraido WHERE campo = 'latitude'")).one()
        assert ocr == antes["valor_bruto"] and float(conf) == antes["confianca"]  # leitura original preservada
        acuracia = {r.campo: r for r in c.execute(text("SELECT * FROM ocr.vw_acuracia_campo"))}
        assert acuracia["latitude"].corrigidos == 1 and acuracia["latitude"].pct_acerto == 0
        assert acuracia["municipio"].corrigidos == 0 and acuracia["municipio"].pct_acerto == 100
        assert acuracia["localidade"].corrigidos == 1
        faixas = c.execute(text("SELECT sum(conferidos) FROM ocr.vw_calibracao_confianca")).scalar()
        assert faixas > 0


def test_ct99_filtros_da_lista(cliente):
    from datetime import datetime, timedelta
    from zoneinfo import ZoneInfo

    teresina = "-42°50'00\""  # dentro do polígono fictício de Teresina no stub
    _enviar(cliente, F.pdf_requerimento(F.DadosReq()), "a.pdf")  # Maria das Graças Sousa, Altos, 01/10/2026
    _enviar(cliente, F.pdf_requerimento(F.DadosReq(
        nome="José Pereira Lima", procedimento="Outorga de Direito de Uso", municipio="TERESINA",
        longitude=teresina, data="15/09/2026")), "b.pdf")
    _enviar(cliente, F.pdf_requerimento(F.DadosReq(
        nome="João da Silva", municipio="TERESINA", longitude=teresina, data="03/10/2026")), "c.pdf")

    def nomes(**params):
        r = cliente.get("/documentos", params=params)
        assert r.status_code == 200, r.text
        return sorted(d["nome_original"] for d in r.json()["itens"])

    assert nomes(requerente="jose") == ["b.pdf"]  # sem acento e sem diferenciar maiúsculas
    assert nomes(requerente="GRAÇAS") == ["a.pdf"]
    assert nomes(requerente="%") == []  # curinga do LIKE é texto comum
    assert nomes(procedimento="Outorga de Direito de Uso") == ["b.pdf"]
    assert nomes(municipio="Teresina") == ["b.pdf", "c.pdf"]
    assert nomes(municipio="Teresina", requerente="joão") == ["c.pdf"]
    assert nomes(data_de="2026-10-01") == ["a.pdf", "c.pdf"]
    assert nomes(data_ate="2026-09-30") == ["b.pdf"]
    assert nomes(data_de="2026-10-01", data_ate="2026-10-01") == ["a.pdf"]
    hoje = datetime.now(ZoneInfo("America/Fortaleza")).date()
    assert nomes(campo_data="gravacao", data_de=hoje.isoformat()) == ["a.pdf", "b.pdf", "c.pdf"]
    assert nomes(campo_data="gravacao", data_ate=(hoje - timedelta(days=1)).isoformat()) == []
    assert cliente.get("/documentos", params={"campo_data": "outro"}).status_code == 422

    opcoes = cliente.get("/filtros").json()
    assert {"nome": "Teresina", "total": 2} in opcoes["municipios"] and {"nome": "Altos", "total": 1} in opcoes["municipios"]
    assert {"nome": "Outorga de Direito de Uso", "total": 1} in opcoes["procedimentos"]
