"""Exportação de documentos de um lote em CSV (separador ';', compatível com Excel pt-BR) ou XLSX."""
from __future__ import annotations

import csv
import io
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .extracao.finalidades import POR_CHAVE
from .modelos import Documento

COLUNAS = [
    ("documento", "Documento (uuid)"), ("arquivo", "Arquivo"), ("status", "Status"),
    ("requerente", "Requerente"), ("cpf_cnpj", "CPF/CNPJ"), ("procedimento", "Procedimento"),
    ("intervencao", "Tipo de intervenção"), ("municipio", "Município"), ("localidade", "Localidade"),
    ("latitude", "Latitude"), ("longitude", "Longitude"), ("finalidades", "Finalidades"),
    ("area_total_ha", "Área irrigada total (ha)"), ("detalhes", "Detalhes das finalidades"),
    ("assinatura", "Assinatura"), ("data", "Data do documento"),
    ("confianca", "Confiança (%)"), ("pendencias", "Pendências"),
]


def linha(doc: Documento) -> dict:
    r = doc.requerimento
    base = {"documento": str(doc.uuid), "arquivo": doc.nome_original, "status": doc.status,
            "assinatura": doc.assinatura_tipo or ""}
    if r is None:
        return {**base, "pendencias": doc.erro or ""}
    return {
        **base,
        "requerente": r.requerente_nome or "", "cpf_cnpj": r.cpf_cnpj or "",
        "procedimento": r.procedimento_texto or "", "intervencao": r.tipo_intervencao_texto or "",
        "municipio": r.municipio_texto or "", "localidade": r.localidade or "",
        "latitude": float(r.latitude) if r.latitude is not None else "",
        "longitude": float(r.longitude) if r.longitude is not None else "",
        "finalidades": ", ".join(f.finalidade_texto or "" for f in r.finalidades if f.valido),
        "area_total_ha": float(sum((i.area_ha or 0) for i in r.irrigacao)) if r.irrigacao else "",
        "detalhes": " | ".join(
            f"{f.finalidade_texto}: {resumo_detalhes(f.chave, f.detalhes)}"
            for f in r.finalidades if f.detalhes and f.chave != "irrigacao" and resumo_detalhes(f.chave, f.detalhes)
        ),
        "data": r.data_documento.strftime("%d/%m/%Y") if r.data_documento else "",
        "confianca": float(r.confianca_geral) if r.confianca_geral is not None else "",
        "pendencias": " | ".join(r.pendencias or []),
    }


def para_csv(docs: Iterable[Documento]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow([rot for _, rot in COLUNAS])
    for d in docs:
        dados = linha(d)
        w.writerow([_fmt_br(dados.get(k, "")) for k, _ in COLUNAS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _fmt_br(v) -> str:
    if isinstance(v, float):
        return f"{v:.7f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(v)


def para_xlsx(docs: Iterable[Documento]) -> bytes:
    docs = list(docs)
    wb = Workbook()
    ws = wb.active
    ws.title = "Requerimentos"
    ws.append([rot for _, rot in COLUNAS])
    for d in docs:
        dados = linha(d)
        ws.append([dados.get(k, "") for k, _ in COLUNAS])
    _formatar(ws)

    wi = wb.create_sheet("Irrigação")
    wi.append(["Documento (uuid)", "Requerente", "Linha", "Sistema de irrigação", "Área (ha)", "Cultura", "Confiança (%)"])
    for d in docs:
        if d.requerimento:
            for it in d.requerimento.irrigacao:
                wi.append([str(d.uuid), d.requerimento.requerente_nome or "", it.ordem, it.sistema_texto,
                           float(it.area_ha) if it.area_ha is not None else None, it.cultura_texto,
                           float(it.confianca) if it.confianca is not None else None])
    _formatar(wi)

    wd = wb.create_sheet("Detalhes das finalidades")
    wd.append(["Documento (uuid)", "Requerente", "Finalidade", "Item", "Campo", "Valor"])
    for d in docs:
        if not d.requerimento:
            continue
        for f in d.requerimento.finalidades:
            if f.chave == "irrigacao":
                continue  # já na aba Irrigação
            for item, campo, valor in linhas_detalhes(f.chave, f.detalhes):
                wd.append([str(d.uuid), d.requerimento.requerente_nome or "", f.finalidade_texto, item, campo, valor])
    _formatar(wd)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _formatar(ws) -> None:
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, start=1):
        largura = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)


def linhas_detalhes(chave: str | None, detalhes: dict | None) -> list[tuple[int | None, str, object]]:
    """(item, rótulo, valor) de cada dado próprio da finalidade, na ordem do formulário."""
    fdef = POR_CHAVE.get(chave or "")
    if fdef is None or not detalhes:
        return []
    saida: list[tuple[int | None, str, object]] = [
        (None, c.rotulo, detalhes.get(c.chave)) for c in fdef.campos if detalhes.get(c.chave) is not None
    ]
    for n, item in enumerate(detalhes.get("itens") or [], start=1):
        saida += [(n, c.rotulo, item.get(c.chave)) for c in fdef.colunas if item.get(c.chave) is not None]
    return saida


def resumo_detalhes(chave: str | None, detalhes: dict | None) -> str:
    partes = []
    for item, rotulo, valor in linhas_detalhes(chave, detalhes):
        v = _fmt_br(float(valor)) if isinstance(valor, (int, float)) else str(valor)
        partes.append(f"{'#' + str(item) + ' ' if item else ''}{rotulo}: {v}")
    return "; ".join(partes)
